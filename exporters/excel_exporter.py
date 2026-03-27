from typing import Any, Dict, List

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelExporter:
    def __init__(self, local_weights_by_path: Dict[str, Any], global_weights: Dict[str, float]):
        self.local_weights_by_path = local_weights_by_path
        self.global_weights = global_weights

    def _add_table_header(self, ws, title, col_span):
        ws.append([title])

        row = ws.max_row
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=col_span)
        ws[f"A{row}"].font = Font(bold=True)
        ws[f"A{row}"].alignment = Alignment(horizontal="center")

    def _add_pairwise_comparison_matrix(self, ws, items: List[str], matrix: np.ndarray):
        self._add_table_header(ws, "Pairwise comparison matrix", len(items) + 1)

        ws.append([""] + items)
        for i, row_item in enumerate(items):
            row = [row_item] + [matrix[i, j] for j in range(len(items))]
            ws.append(row)

        for col_idx in range(len(items) + 1):
            col_letter = get_column_letter(col_idx + 1)
            ws.column_dimensions[col_letter].width = 14

        ws.append([])

    def _add_local_weights_matrix(self, ws, data):
        self._add_table_header(ws, "Local weights", 2)

        ws.append(["Element", "Local weight"])
        for item, w in zip(data["items"], data["weights"]):
            ws.append([item, w])
        ws.append([])
        ws.append(["λ max", data["lam_max"]])
        ws.append(["CI", data["ci"]])
        ws.append(["CR", data["cr"]])
        if data["cr"] > 0.1:
            ws.cell(ws.max_row, 2).fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

    def save_results_to_excel_and_csv(self, excel_path: str, csv_path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)

        ws = wb.create_sheet(title="Global weights")
        ws.append(["Risk", "Global weight"])
        for name, weight in sorted(self.global_weights.items(), key=lambda x: -x[1]):
            ws.append([name, weight])

        for path_str, data in self.local_weights_by_path.items():
            safe_title = path_str.replace("/", "⧸")[:31] if path_str else "Stages"
            ws = wb.create_sheet(title=safe_title)

            self._add_pairwise_comparison_matrix(ws, data["items"], data["matrix"])
            self._add_local_weights_matrix(ws, data)

        wb.save(excel_path)

        df = pd.DataFrame(self.global_weights.items(), columns=["Path", "Global Weight"])
        df.to_csv(csv_path, index=False)
